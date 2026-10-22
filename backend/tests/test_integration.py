import asyncio
import hashlib
import hmac
import io
import json
import os
import sys
import time
from pathlib import Path
from urllib.parse import urlencode
from uuid import uuid4
from docx import Document
from pptx import Presentation
import pytest
from app.core.config import settings
from app.services.commerce import rpc
from app.utils.telegram_auth import validate_telegram_init_data, create_access_token


def signed_init(user_id, name='A+B & 50% Test', auth_date=None, start_param=None):
    fields = {'auth_date': str(int(time.time()) if auth_date is None else auth_date),
        'user': json.dumps({'id': user_id, 'first_name': name, 'is_bot': False}, ensure_ascii=False)}
    if start_param:
        fields['start_param'] = start_param
    secret = hmac.new(b'WebAppData', settings.TELEGRAM_BOT_TOKEN.encode(), hashlib.sha256).digest()
    fields['hash'] = hmac.new(secret, '\n'.join(f'{key}={value}' for key, value in sorted(fields.items())).encode(), hashlib.sha256).hexdigest()
    return urlencode(fields)


def login(client, user_id):
    response = client.post('/api/v1/auth/telegram', json={'init_data': signed_init(user_id)})
    assert response.status_code == 200, response.text
    return {'Authorization': 'Bearer ' + response.json()['access_token']}


def completed_job(client, headers, job_id):
    deadline = time.monotonic() + 20
    while time.monotonic() < deadline:
        result = client.get('/api/v1/jobs/' + job_id, headers=headers)
        assert result.status_code == 200
        if result.json()['status'] in {'completed', 'failed'}:
            return result.json()
        time.sleep(.1)
    pytest.fail('Generation job did not finish')


def test_telegram_signature_and_no_dev_backdoor(client):
    raw = signed_init(88001)
    assert validate_telegram_init_data(raw)['first_name'] == 'A+B & 50% Test'
    assert validate_telegram_init_data(signed_init(88001, auth_date='broken')) is None
    assert validate_telegram_init_data(signed_init(88001, auth_date=int(time.time()) + 100)) is None
    assert validate_telegram_init_data(raw.replace('50%25', '51%25')) is None
    assert client.post('/api/v1/auth/dev-login').status_code == 404
    assert client.post('/api/v1/bot/set-webhook').status_code == 404
    assert client.post('/api/v1/bot/webhook', json={'message': {}}).status_code == 404
    assert client.get('/api/v1/commerce').status_code in {401, 403}


def test_bot_to_api_shared_referrals_and_real_files(client, monkeypatch):
    # Both repositories use the same PostgREST/Supabase client against the test PostgreSQL engine.
    bot_root = Path(os.environ.get('ORZU_BOT_PATH', str(Path(__file__).resolve().parents[2].parent / 'orzu-bot')))
    if not bot_root.is_dir():
        bot_root = Path(__file__).resolve().parents[2].parent / 'Orzu'
    if not bot_root.is_dir():
        pytest.skip('Check out Orzu next to this repository to run the cross-repository test')
    monkeypatch.setenv('BOT_TOKEN', settings.TELEGRAM_BOT_TOKEN)
    sys.path.insert(0, str(bot_root))
    import database as bot_database
    import main as bot_main
    monkeypatch.setattr(bot_database, '_client', client.db)

    class Person:
        is_bot = False
        def __init__(self, user_id):
            self.id, self.first_name = user_id, 'Student'
        def model_dump(self):
            return {'id': self.id, 'first_name': self.first_name, 'is_bot': False}

    class FakeMessage:
        def __init__(self, user):
            self.from_user = user
            self.answers = []
        async def answer(self, text, **kwargs):
            self.answers.append((text, kwargs))
            return self

    class State:
        def __init__(self): self.data = {}
        async def clear(self): self.data = {}
        async def update_data(self, **kwargs): self.data.update(kwargs)

    class Bot:
        async def get_chat_member(self, *args):
            return type('Member', (), {'status': 'member'})()
        async def send_message(self, *args, **kwargs): pass

    async def onboard():
        from aiogram.filters import CommandObject
        bot = Bot()
        for user_id in [89000, 89001, 89002, 89003, 89004, 89005]:
            message, state = FakeMessage(Person(user_id)), State()
            await bot_main.start_handler_local(message, bot, state,
                CommandObject(command='start', args=None if user_id == 89000 else 'ref_89000'))
            assert message.answers
        # Replaying a referral link cannot award a second batch.
        await bot_main.confirm_referral(89005, bot)
    asyncio.run(onboard())
    headers = login(client, 89000)
    account = client.get('/api/v1/commerce', headers=headers).json()
    assert account['confirmed_referrals'] == 5
    assert account['referral_presentations'] == 5
    assert account['course_credits'] == 2
    assert account['referral_credits'] == 5
    assert account['access_mode'] == 'referral_only'
    assert 'payment_card' not in account
    assert account['referral_link'].endswith('?start=ref_89000')
    assert not any(key.startswith('is_pro') for key in account)
    request_id = str(uuid4())
    form = {'topic': 'History of the Silk Road', 'author': 'Test Author', 'language': 'uz', 'subject': 'minimalizm', 'variant': '1'}
    headers['Idempotency-Key'] = request_id
    response = client.post('/api/v1/presentations/generate', headers=headers, data=form)
    assert response.status_code == 202, response.text
    duplicate = client.post('/api/v1/presentations/generate', headers=headers, data=form)
    assert duplicate.status_code == 202
    job = completed_job(client, headers, request_id)
    assert job['status'] == 'completed', job
    assert job['telegram_sent'] is True
    download = client.get(job['download_url'], headers=headers)
    assert download.status_code == 200
    deck = Presentation(io.BytesIO(download.content))
    assert len(deck.slides) == 20
    assert not any('{{' in shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert client.get('/api/v1/commerce', headers=headers).json()['referral_presentations'] == 4
    stranger = login(client, 89006)
    assert client.get('/api/v1/jobs/' + request_id, headers=stranger).status_code == 404
    assert client.get(job['download_url'], headers=stranger).status_code == 404
    changed = client.post('/api/v1/presentations/generate', headers=headers, data={**form, 'topic': 'Different topic'})
    assert changed.status_code == 409

    headers['Idempotency-Key'] = str(uuid4())
    response = client.post('/api/v1/articles/generate', headers=headers,
        data={'topic': 'History of teaching', 'author': 'Test Author', 'language': 'en'})
    assert response.status_code == 202, response.text
    job = completed_job(client, headers, response.json()['id'])
    assert job['status'] == 'completed', job
    word = client.get(job['download_url'], headers=headers)
    assert word.status_code == 200
    document = Document(io.BytesIO(word.content))
    assert 'Methodology' in '\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert client.get('/api/v1/commerce', headers=headers).json()['referral_credits'] == 4
    assert asyncio.run(bot_database.get_user_balance(89000)) == 0


def test_initial_access_subscription_and_invalid_inputs(client):
    headers = login(client, 90000)
    request_id = str(uuid4())
    headers['Idempotency-Key'] = request_id
    form = {'topic': 'History of education', 'author': 'Student', 'subject': 'minimalizm'}
    response = client.post('/api/v1/presentations/generate', headers=headers, data=form)
    assert response.status_code == 402
    assert response.json()['code'] == 'insufficient_referral_credits'
    assert client.post('/api/v1/commerce/subscribe', headers=headers).status_code == 410
    client.db.table('users').update({'balance': 35000}).eq('telegram_id', 90000).execute()
    assert client.post('/api/v1/commerce/subscribe', headers=headers).status_code == 410
    assert client.post('/api/v1/commerce/subscribe', headers=headers).status_code == 410
    account = client.get('/api/v1/commerce', headers=headers).json()
    assert account['balance'] == 35000
    assert account['referral_credits'] == 0
    assert account['subscription_remaining'] == 0
    assert client.post('/api/v1/presentations/generate', headers=headers, data={**form, 'subject': '../tarix'}).status_code == 400
    assert client.post('/api/v1/presentations/generate', headers=headers, data={**form, 'language': 'xx'}).status_code == 422
    assert client.post('/api/v1/presentations/generate', headers=headers, data=form,
        files={'images': ('file.png', b'not an image', 'image/png')}).status_code == 400
    assert client.get('/api/v1/commerce', headers=headers).json()['referral_credits'] == 0


def test_generation_failure_refunds_allowance(client, monkeypatch):
    headers = login(client, 91000)
    friend = 91001
    client.db.rpc('commerce_register', {'p_user_id': friend, 'p_referrer_id': 91000}).execute()
    client.db.rpc('commerce_confirm_referral', {'p_user_id': friend}).execute()
    from app.services.presentation.ai_generator import AIContentGenerator
    async def failed(self, **kwargs): raise RuntimeError('Provider unavailable')
    monkeypatch.setattr(AIContentGenerator, 'generate_template_slides', failed)
    headers['Idempotency-Key'] = str(uuid4())
    response = client.post('/api/v1/presentations/generate', headers=headers,
        data={'topic': 'History of education', 'author': 'Student'})
    assert response.status_code == 202, response.text
    job = completed_job(client, headers, response.json()['id'])
    assert job['status'] == 'failed'
    assert client.get('/api/v1/commerce', headers=headers).json()['referral_credits'] == 1
    assert client.get('/api/v1/jobs/' + response.json()['id'] + '/download', headers=headers).status_code == 409
