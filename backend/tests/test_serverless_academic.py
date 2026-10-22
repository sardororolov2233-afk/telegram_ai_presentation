from pathlib import Path
import pytest
from docx import Document
from app.services.academic.pipeline import AcademicPipeline, generate_referat, AcademicInput
from app.core.config import settings


@pytest.mark.asyncio
@pytest.mark.parametrize('service', ['course', 'referat'])
async def test_academic_pipeline_renders_real_docx_with_original_cover(service, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, 'JOB_STORAGE_DIR', str(tmp_path))
    calls = []

    async def fake_ai(prompt):
        calls.append(prompt)
        return 'Academic analysis with enough detail to form a paragraph. ' * 40

    async def delivered(*args):
        return True

    monkeypatch.setattr('app.services.academic.pipeline.ask_ai', fake_ai)
    monkeypatch.setattr('app.services.academic.course.ai_request', fake_ai)
    monkeypatch.setattr('app.services.academic.pipeline.send_document', delivered)
    plan = 'REJA\nKIRISH\nI BOB: Kushonlar tarixi\n1.1. Nazariy asoslar\n1.2. Tarix\nII BOB: Amaliyot\n2.1. Tahlil\n2.2. Natijalar\nIII BOB: Takliflar\nXULOSA\nADABIYOTLAR'
    result = await AcademicPipeline(service).run(telegram_id=123, tema='Kushonlar tarixi', plan=plan,
        muallif='Sardor Orolov', universitet='Test University', fakultet='History', kurs_guruh='4-kurs',
        sahifa=30, ish_turi='KURS ISHI' if service == 'course' else 'REFERAT')
    path = Path(result['document_path'])
    assert path.is_relative_to(tmp_path / 'outputs')
    assert path.is_file() and result['telegram_sent']
    doc = Document(path)
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Sardor Orolov' in text
    assert 'TEST UNIVERSITY' in text
    assert 'Kushonlar tarixi' in text
    assert len(calls) >= 6
    assert 'Academic analysis' in text
    assert doc.sections[0].left_margin.mm == pytest.approx(30, abs=0.01)


@pytest.mark.asyncio
async def test_failed_delivery_retains_file_for_owned_job_download(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, 'JOB_STORAGE_DIR', str(tmp_path))

    async def fake_text(*args):
        return 'KIRISH\nReliable academic text.\nXULOSA\nFindings.'

    async def not_delivered(*args):
        return False

    monkeypatch.setattr('app.services.academic.pipeline.generate_referat', fake_text)
    monkeypatch.setattr('app.services.academic.pipeline.send_document', not_delivered)
    result = await AcademicPipeline('referat').run(telegram_id=123, tema='History topic',
        plan='REJA\nKIRISH\nI BOB\nXULOSA', muallif='Test Author')
    assert Path(result['document_path']).is_file()
    assert not result['telegram_sent']


@pytest.mark.asyncio
async def test_generation_failure_cannot_leave_a_partial_file(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, 'JOB_STORAGE_DIR', str(tmp_path))

    async def fail(*args):
        raise RuntimeError('AI unavailable')

    monkeypatch.setattr('app.services.academic.pipeline.generate_referat', fail)
    with pytest.raises(RuntimeError):
        await AcademicPipeline('referat').run(telegram_id=123, tema='History topic',
            plan='REJA\nKIRISH\nI BOB\nXULOSA', muallif='Test Author')
    assert not list(tmp_path.rglob('*.docx'))


@pytest.mark.parametrize('service, units', [('course', 2), ('referat', 0)])
def test_referral_only_queue_worker_renders_and_charges_once(client, monkeypatch, service, units):
    import io
    from uuid import uuid4
    from test_integration import login, completed_job

    async def text(*args, **kwargs):
        return 'Academic analysis with concrete findings. ' * 40

    async def delivered(*args, **kwargs):
        return True

    monkeypatch.setattr('app.services.academic.pipeline.ask_ai', text)
    monkeypatch.setattr('app.services.academic.course.ai_request', text)
    monkeypatch.setattr('app.services.academic.pipeline.send_document', delivered)
    uid = 93200 if service == 'course' else 93300
    headers = login(client, uid)
    for friend in [uid + 1, uid + 2]:
        client.db.rpc('commerce_register', {'p_user_id':friend,'p_first_name':'Friend','p_referrer_id':uid}).execute()
        client.db.rpc('commerce_confirm_referral', {'p_user_id':friend}).execute()
    payload = {'tema':'History research','muallif':'Test Author','universitet':'University',
        'plan':'REJA\nKIRISH\nI BOB. History\nII BOB. Analysis\nXULOSA\nADABIYOTLAR',
        'ish_turi':'KURS ISHI' if service == 'course' else 'REFERAT', 'send_to_telegram':True}
    request = str(uuid4())
    args = {'p_user_id':uid,'p_service':service,'p_request_id':request,'p_payload':payload}
    charge = client.db.rpc('commerce_enqueue',args).execute().data
    assert charge['credit_units'] == units and charge['amount'] == 0
    assert not client.db.rpc('commerce_enqueue',args).execute().data['created']
    job = completed_job(client,headers,request)
    assert job['status'] == 'completed' and job['telegram_sent']
    download = client.get(job['download_url'],headers=headers)
    assert download.status_code == 200
    doc = Document(io.BytesIO(download.content))
    assert any('Test Author' in p.text for p in doc.paragraphs)
    account = client.get('/api/v1/commerce',headers=headers).json()
    assert account['referral_credits'] == 2 - units
    health = client.get('/health').json()
    assert 'serverless-referral-only-v2' in health['features']
