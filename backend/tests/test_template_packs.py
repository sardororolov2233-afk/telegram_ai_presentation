import pytest
from pptx import Presentation
from app.services.presentation.subject_templates import list_subject_templates, load_manifest, build_subject_schema
from app.services.presentation.pptx_generator import generate_template_pptx


def test_all_supplied_subjects_are_in_one_catalog():
    catalog = list_subject_templates()
    assert len(catalog) == 17
    assert sum(len(subject['variants']) for subject in catalog) == 40
    for subject in catalog:
        for variant in subject['variants']:
            schema = build_subject_schema(subject['subject'], variant['variant'], 'Education', 'Sardor')
            assert schema
            assert variant['slide_count'] > 0


@pytest.mark.parametrize('subject', ['biologiya', 'adabiyot', 'bank'])
async def test_existing_legacy_pack_still_produces_a_filled_deck(subject, tmp_path):
    manifest = load_manifest(subject, 1)
    data = {key: 'Matn' for key in manifest['fields']}
    data['muallif'] = 'Sardor'
    path, count = await generate_template_pptx(data, str(tmp_path / f'{subject}.pptx'), subject=subject)
    presentation = Presentation(path)
    assert len(presentation.slides) == count == manifest['slide_count']
    assert not any('{{' in shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
