"""Claude Haiku writes structured text. Trusted Python code renders the DOCX."""
import asyncio
import json
import re
from pathlib import Path
from typing import Optional
from uuid import uuid4
import httpx
from pydantic import BaseModel, Field
from docx import Document
from docx.shared import Cm, Pt
from app.core.config import settings
from app.services.presentation.telegram_sender import send_document


class ArticleSection(BaseModel):
    heading: str = Field(min_length=1)
    paragraphs: list[str] = Field(min_length=1)


class ArticleContent(BaseModel):
    title: str = Field(min_length=1)
    abstract: str = Field(min_length=50)
    keywords: list[str] = Field(min_length=3)
    sections: list[ArticleSection] = Field(min_length=3, max_length=10)
    conclusion: str = Field(min_length=50)
    references: list[str] = Field(default_factory=list)


class ArticlePipeline:
    model = settings.ACADEMIC_MODEL

    async def _call_llm(self, system_prompt, user_prompt):
        if not settings.OPENROUTER_API_KEY:
            raise RuntimeError('OPENROUTER_API_KEY sozlanmagan')
        async with httpx.AsyncClient(timeout=240.0) as client:
            response = await client.post('https://openrouter.ai/api/v1/chat/completions',
                headers={'Authorization': f'Bearer {settings.OPENROUTER_API_KEY}'}, json={
                    'model': self.model, 'messages': [{'role': 'system', 'content': system_prompt},
                        {'role': 'user', 'content': user_prompt}], 'temperature': 0.3, 'max_tokens': 16384})
            response.raise_for_status()
        choice = response.json()['choices'][0]
        if choice.get('finish_reason') == 'length':
            raise ValueError('Maqola matni tugallanmagan')
        text = re.sub(r'<think>.*?</think>', '', choice['message'].get('content') or '', flags=re.S).strip()
        return re.sub(r'^```(?:json)?\s*|\s*```$', '', text).strip()

    @staticmethod
    def write_document(content, author, language, path):
        labels = {
            'uz': ('Annotatsiya', 'Kalit so‘zlar', 'Xulosa', 'Foydalanilgan adabiyotlar', 'Manbalar taqdim etilmagan.'),
            'ru': ('Аннотация', 'Ключевые слова', 'Заключение', 'Литература', 'Источники не предоставлены.'),
            'en': ('Abstract', 'Keywords', 'Conclusion', 'References', 'No sources were supplied.'),
        }[language]
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin, section.right_margin = Cm(3), Cm(1.5)
        normal = doc.styles['Normal']
        normal.font.name, normal.font.size = 'Times New Roman', Pt(14)
        normal.paragraph_format.line_spacing = 1.5
        doc.add_heading(content.title, 0)
        doc.add_paragraph(author)
        doc.add_heading(labels[0], 1)
        doc.add_paragraph(content.abstract)
        doc.add_paragraph(labels[1] + ': ' + ', '.join(content.keywords))
        for part in content.sections:
            doc.add_heading(part.heading, 1)
            for text in part.paragraphs:
                doc.add_paragraph(text)
        doc.add_heading(labels[2], 1)
        doc.add_paragraph(content.conclusion)
        doc.add_heading(labels[3], 1)
        for reference in content.references or [labels[4]]:
            doc.add_paragraph(reference)
        doc.save(path)

    async def run(self, topic: str, author: str, language: str, telegram_id: Optional[int], references: str = ''):
        languages = {'uz': 'Uzbek (Latin)', 'ru': 'Russian', 'en': 'English'}
        prompt = (f'Write a complete academic draft in {languages[language]}, topic: {json.dumps(topic)}. '
            'Include an introduction, methodology, discussion and conclusion. Target 1500–2200 words. '
            'Distinguish evidence from inference; never invent research results, quotations, bibliographic details, '
            'UDC codes or claims of journal acceptance. Use only supplied sources in references. '
            f'Supplied sources: {json.dumps(references)}. With no sources use an empty references list. '
            f'Return only JSON matching this schema: {json.dumps(ArticleContent.model_json_schema())}')
        raw = await self._call_llm('You write academic text, never executable code. Return a JSON object.', prompt)
        content = ArticleContent.model_validate_json(raw)
        directory = Path(settings.JOB_STORAGE_DIR) / 'outputs'
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f'article_{uuid4()}.docx'
        try:
            await asyncio.to_thread(self.write_document, content, author, language, str(path))
            sent = await send_document(telegram_id, str(path), f'📝 {topic[:900]}') if telegram_id else False
            return {'document_path': str(path), 'telegram_sent': sent}
        except BaseException:
            path.unlink(missing_ok=True)
            raise
